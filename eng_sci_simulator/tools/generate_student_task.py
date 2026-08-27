from pathlib import Path
import math
import os
import sys

sys.path.insert(0, "/private/tmp/quakelab_doc_deps")

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "student-resources"
ASSETS = OUT / "assets"
DOCX = OUT / "Year_9_Earthquake_Engineering_Student_Task.docx"

NAVY = "18324A"
BLUE = "1F5D7A"
TEAL = "2A7F7A"
PALE_BLUE = "E8EEF5"
PALE_TEAL = "E8F4F2"
PALE_GOLD = "FFF2CC"
PALE_RED = "FCE8E6"
DARK = "24323D"
MID = "50616F"
WHITE = "FFFFFF"
GRID = "B8C5CE"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for item in candidates:
        if Path(item).exists():
            try:
                return ImageFont.truetype(item, size)
            except OSError:
                pass
    return ImageFont.load_default()


def arrow(draw, start, end, fill, width=7, head=16):
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    p1 = (end[0] - head * math.cos(angle - 0.55), end[1] - head * math.sin(angle - 0.55))
    p2 = (end[0] - head * math.cos(angle + 0.55), end[1] - head * math.sin(angle + 0.55))
    draw.polygon([end, p1, p2], fill=fill)


def label(draw, xy, text, size=26, color=(36, 50, 61), bold=False, anchor=None):
    draw.text(xy, text, fill=color, font=font(size, bold), anchor=anchor)


def save_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    navy = (24, 50, 74)
    blue = (31, 93, 122)
    teal = (42, 127, 122)
    gold = (229, 164, 48)
    red = (190, 72, 65)
    grey = (90, 108, 120)
    pale = (242, 246, 248)

    # 1. Ground movement, inertia and load path.
    img = Image.new("RGB", (1400, 680), "white")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((35, 35, 1365, 645), 28, fill=pale, outline=(210, 220, 226), width=3)
    d.line((110, 550, 1290, 550), fill=navy, width=18)
    d.rectangle((510, 180, 890, 550), outline=navy, width=12)
    d.line((510, 360, 890, 360), fill=navy, width=10)
    d.line((510, 550, 890, 180), fill=teal, width=14)
    d.line((890, 550, 510, 180), fill=teal, width=14)
    for x, y in [(510, 550), (890, 550), (510, 360), (890, 360), (510, 180), (890, 180)]:
        d.ellipse((x-13, y-13, x+13, y+13), fill="white", outline=navy, width=5)
    arrow(d, (180, 610), (430, 610), gold, 12, 28)
    arrow(d, (1040, 130), (820, 130), red, 12, 28)
    arrow(d, (700, 215), (700, 520), blue, 8, 22)
    label(d, (180, 575), "GROUND MOVEMENT", 25, navy, True)
    label(d, (1040, 85), "INERTIA", 25, red, True)
    label(d, (730, 360), "load path to\nfoundation", 24, blue, True)
    label(d, (80, 75), "Earthquake forces need a continuous route through", 33, navy, True)
    label(d, (80, 120), "floors, joints, bracing and foundations.", 33, navy, True)
    img.save(ASSETS / "forces_load_path.png")

    # 2. Racking and triangulation.
    img = Image.new("RGB", (1400, 620), "white")
    d = ImageDraw.Draw(img)
    panels = [(45, 60, 435, 555), (505, 60, 895, 555), (965, 60, 1355, 555)]
    titles = ["UNBRACED", "ONE DIAGONAL", "X-BRACED"]
    for box, title in zip(panels, titles):
        d.rounded_rectangle(box, 22, fill=pale, outline=(210, 220, 226), width=3)
        label(d, ((box[0]+box[2])//2, 105), title, 27, navy, True, "mm")
    # distorted frame
    d.line([(120, 470), (315, 470), (365, 210), (170, 210), (120, 470)], fill=red, width=13)
    arrow(d, (145, 175), (350, 175), red, 9, 23)
    label(d, (240, 520), "racks sideways", 24, red, True, "mm")
    # one diagonal
    d.rectangle((580, 210, 820, 470), outline=navy, width=13)
    d.line((580, 470, 820, 210), fill=teal, width=15)
    arrow(d, (595, 175), (800, 175), gold, 9, 23)
    label(d, (700, 520), "triangle resists shape change", 22, teal, True, "mm")
    # X
    d.rectangle((1040, 210, 1280, 470), outline=navy, width=13)
    d.line((1040, 470, 1280, 210), fill=teal, width=15)
    d.line((1280, 470, 1040, 210), fill=teal, width=15)
    arrow(d, (1055, 175), (1260, 175), gold, 9, 23)
    label(d, (1160, 520), "two diagonal load paths", 22, teal, True, "mm")
    img.save(ASSETS / "triangulation.png")

    # 3. Three protection strategies.
    img = Image.new("RGB", (1400, 700), "white")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((35, 35, 1365, 665), 28, fill=pale, outline=(210, 220, 226), width=3)
    centers = [250, 700, 1150]
    titles = ["BRACING", "BASE ISOLATION", "DAMPING"]
    subtitles = ["stiffens the frame", "reduces force transfer", "dissipates energy"]
    for cx, title, sub in zip(centers, titles, subtitles):
        label(d, (cx, 95), title, 28, navy, True, "mm")
        label(d, (cx, 135), sub, 22, grey, False, "mm")
        d.line((cx-150, 555, cx+150, 555), fill=navy, width=14)
        d.rectangle((cx-115, 220, cx+115, 535), outline=navy, width=11)
        d.line((cx-115, 375, cx+115, 375), fill=navy, width=8)
    d.line((135, 535, 365, 220), fill=teal, width=14)
    d.line((365, 535, 135, 220), fill=teal, width=14)
    for x in [625, 680, 735, 790]:
        d.ellipse((x-22, 535, x+22, 570), fill=gold, outline=navy, width=3)
    d.line((1035, 535, 1265, 220), fill=teal, width=11)
    d.rounded_rectangle((1105, 330, 1195, 430), 12, fill=(255, 242, 204), outline=gold, width=6)
    label(d, (700, 620), "No single strategy makes a building ‘earthquake-proof’. Engineers combine systems.", 27, navy, True, "mm")
    img.save(ASSETS / "protection_strategies.png")

    # 4. Resonance response curve.
    img = Image.new("RGB", (1400, 700), "white")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((35, 35, 1365, 665), 28, fill=pale, outline=(210, 220, 226), width=3)
    x0, y0, x1, y1 = 150, 560, 1270, 150
    d.line((x0, y0, x1, y0), fill=navy, width=7)
    d.line((x0, y0, x0, y1), fill=navy, width=7)
    arrow(d, (x1-30, y0), (x1+20, y0), navy, 7, 18)
    arrow(d, (x0, y1+30), (x0, y1-20), navy, 7, 18)
    label(d, (710, 620), "shaking frequency (Hz)", 26, navy, True, "mm")
    label(d, (65, 355), "response", 25, navy, True, "mm")
    pts1, pts2 = [], []
    for i in range(501):
        x = x0 + (x1-x0)*i/500
        f = 0.2 + 4.8*i/500
        r1 = 0.18 + 1.0/(1+((f-1.15)/0.22)**2)
        r2 = 0.16 + 0.9/(1+((f-3.25)/0.35)**2)
        pts1.append((x, y0-300*r1/1.2))
        pts2.append((x, y0-300*r2/1.2))
    d.line(pts1, fill=teal, width=10)
    d.line(pts2, fill=gold, width=10)
    label(d, (380, 205), "taller / more flexible", 23, teal, True, "mm")
    label(d, (925, 245), "shorter / stiffer", 23, (180, 122, 20), True, "mm")
    d.line((405, 250, 405, 555), fill=teal, width=3)
    d.line((880, 285, 880, 555), fill=gold, width=3)
    label(d, (710, 80), "Largest response can occur when frequencies are close", 31, navy, True, "mm")
    img.save(ASSETS / "resonance.png")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=DARK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths=None, font_size=8.5, row_heights=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    t.alignment = 1
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, NAVY, font_size)
        shade(t.rows[0].cells[i], PALE_BLUE)
        t.rows[0].cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(t.rows[0])
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, DARK, font_size)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                shade(cells[i], "F7F9FA")
        if row_heights and ridx < len(row_heights):
            t.rows[-1].height = Inches(row_heights[ridx])
            t.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    for row in t.rows:
        for cell in row.cells:
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, field):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, txt, end])


def set_repeat_table_layout(table_obj):
    tbl_pr = table_obj._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    width = tbl_pr.find(qn("w:tblW"))
    if width is not None:
        width.set(qn("w:w"), "9360")
        width.set(qn("w:type"), "dxa")


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after, color in [
        ("Title", 27, 0, 10, NAVY),
        ("Subtitle", 13, 0, 10, TEAL),
        ("Heading 1", 16, 18, 10, BLUE),
        ("Heading 2", 13, 14, 7, BLUE),
        ("Heading 3", 11.5, 10, 5, NAVY),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    return doc


def header_footer(doc):
    for sec in doc.sections:
        header = sec.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("QUAKELAB  /  YEAR 9 EARTHQUAKE ENGINEERING")
        r.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(BLUE)
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("DESIGN • TEST • MEASURE • IMPROVE     |     ")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MID)
        add_field(p, "PAGE")


def add_title(doc, kicker, title, subtitle=None):
    p = doc.add_paragraph()
    if getattr(doc, "_quakelab_page_break", False):
        p.paragraph_format.page_break_before = True
        doc._quakelab_page_break = False
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(kicker.upper())
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_heading(title, 0)
    if subtitle:
        p = doc.add_paragraph(style="Subtitle")
        p.add_run(subtitle)


def callout(doc, title, text, fill=PALE_TEAL):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c, fill)
    set_cell_margins(c, 130, 180, 130, 180)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def lines(doc, count=3):
    for _ in range(count):
        p = doc.add_paragraph("________________________________________________________________________________")
        p.paragraph_format.space_after = Pt(3)
        p.runs[0].font.color.rgb = RGBColor.from_string(GRID)
        p.runs[0].font.size = Pt(8)


def question(doc, number, text, lines_count=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}. {text}")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    lines(doc, lines_count)


def page(doc):
    doc._quakelab_page_break = True


def add_source(doc, title, organisation, url, note):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{organisation} — ")
    r.bold = True
    add_hyperlink(p, title, url)
    p = doc.add_paragraph(note)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(6)
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor.from_string(MID)


def build_doc():
    save_diagrams()
    doc = setup_document()

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("QUAKELAB INVESTIGATION")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_heading("Earthquake Engineering", 0)
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Structural basics through experimentation — then an open design challenge")
    doc.add_picture(str(ASSETS / "forces_load_path.png"), width=Inches(6.75))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    table(doc, ["Student", "Class", "Teacher", "Date"], [["", "", "", ""]], [2.7, 1.25, 1.4, 1.25], 9, [0.5])
    callout(doc, "Driving question", "How can evidence about forces, shape, materials, mass and vibration help us design safer buildings?", PALE_GOLD)
    p = doc.add_paragraph("Year 9 Science & Engineering  •  5–7 lessons  •  Individual task")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(BLUE)

    # Overview
    page(doc)
    add_title(doc, "Start here", "Your engineering mission", "You will build knowledge in four controlled investigations before applying it to your own solution.")
    doc.add_heading("Learning intentions", 1)
    for text in [
        "Explain how earthquake shaking creates forces in a structure.",
        "Investigate how bracing, materials, mass position and frequency affect structural response.",
        "Collect sensor data, identify patterns, and make evidence-based claims.",
        "Research earthquake-resistant design principles and judge the credibility of sources.",
        "Design, test and improve a solution within realistic constraints.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Success criteria", 1)
    table(doc, ["I can…", "Evidence to include"], [
        ["conduct a fair test", "one independent variable; other settings controlled"],
        ["use measurements", "units, sensor position, maximum response and cost"],
        ["explain cause and effect", "Claim–Evidence–Reasoning (CER), not just ‘it worked’"],
        ["improve a design", "labelled versions and a justified change"],
        ["use research responsibly", "credible sources, notes in my own words, links recorded"],
    ], [2.2, 4.5], 9)
    callout(doc, "Important model limitation", "QuakeLab is a simplified classroom model. A ‘survived’ result is not proof that a real building is safe or code-compliant. Real engineers use site data, standards, detailed analysis, testing and professional judgement.", PALE_RED)
    doc.add_heading("Simulation survival checks", 2)
    p = doc.add_paragraph("The app reports ‘survived’ only when all four simplified limits are met:")
    p.paragraph_format.space_after = Pt(4)
    table(doc, ["Roof sway", "Storey drift", "Failed members", "Height retained"], [["≤ 160 mm", "≤ 80 mm", "≤ 25%", "≥ 65%"]], [1.65]*4, 9)

    # Research primer
    page(doc)
    add_title(doc, "Research primer", "What happens during an earthquake?", "Read, annotate the image, then answer using scientific language.")
    doc.add_picture(str(ASSETS / "forces_load_path.png"), width=Inches(6.7))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph("Ground motion travels as seismic waves. Buildings resist rapid movement because of inertia, creating changing horizontal forces. Damage depends on the motion’s amplitude, frequency and duration, plus distance, depth, local geology, foundations, structural form and construction quality.")
    p.paragraph_format.space_after = Pt(8)
    callout(doc, "Australian context", "Australia is not on a major plate boundary, but damaging earthquakes do occur. The 1989 Newcastle earthquake caused deaths and major damage. Earthquake hazard information is used in mitigation and building-code decisions.", PALE_GOLD)
    question(doc, "R1", "Distinguish earthquake magnitude from shaking intensity at a particular place.", 2)
    question(doc, "R2", "Use the diagram to explain why a continuous load path from roof to foundation matters.", 2)
    question(doc, "R3", "Name two site factors and two building factors that could change the amount of damage.", 2)

    # Principles
    page(doc)
    add_title(doc, "Structural basics", "Five ideas engineers combine", "A safer design balances performance, cost, materials, use and constructability.")
    doc.add_picture(str(ASSETS / "protection_strategies.png"), width=Inches(6.7))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    table(doc, ["Principle", "Meaning", "Possible strategy"], [
        ["Load path", "Forces travel through connected parts to the ground.", "Strong joints; tied floors/walls; secure foundations"],
        ["Strength + stiffness", "Strength resists failure; stiffness limits shape change and drift.", "Braced frames; shear walls; adequate members"],
        ["Ductility", "A system deforms and carries load instead of failing suddenly.", "Detailed steel/reinforcement; ductile connections"],
        ["Energy control", "Reduce transferred energy or dissipate it.", "Base isolation; dampers; material damping"],
        ["Regularity + mass", "Balanced load paths and sensible mass placement reduce twisting and demand.", "Symmetry; avoid weak storeys; lighter upper levels"],
    ], [1.25, 2.75, 2.7], 8.5)
    question(doc, "R4", "Why is ‘earthquake-resistant’ more accurate than ‘earthquake-proof’?")
    question(doc, "R5", "Explain the difference between strength, stiffness and ductility. Why might maximising only stiffness be a poor strategy?", 3)

    # Simulator quick start
    page(doc)
    add_title(doc, "Before testing", "QuakeLab investigation routine", "Use the same routine for every trial so your data can be compared.")
    table(doc, ["Step", "Action", "Quality check"], [
        ["1", "Build from Floor Joints. Join joints with Beams and/or Braces.", "Is there a continuous path to the shake table?"],
        ["2", "Choose materials. Attach payload in 100 kg intervals.", "Record material, mass and total cost."],
        ["3", "Place up to four Sensors on joints or along members.", "Record exactly where S1–S4 are located."],
        ["4", "Set amplitude, frequency, duration and damping.", "Change only the planned independent variable."],
        ["5", "Start test. Observe motion and failures; inspect response graph.", "Record maximum values with units."],
        ["6", "Retest or modify after the test. Save evidence if required.", "Label each version; do not overwrite observations."],
    ], [0.45, 3.6, 2.65], 8.7)
    callout(doc, "Sensor meaning", "A sensor reports signed horizontal movement (Δx) at its attachment point relative to the shake table. Positive and negative show direction. For comparison, use the maximum absolute value: the largest distance from zero.", PALE_TEAL)
    doc.add_heading("Fair-test record", 2)
    table(doc, ["Independent variable", "Dependent variable(s)", "Controlled variables"], [["what I change", "what I measure", "what I keep the same"]], [2.1, 2.2, 2.4], 9, [0.6])
    question(doc, "Q0", "Why should you reset or use the same starting structure before comparing trials?")

    # Activity 1
    page(doc)
    add_title(doc, "Investigation 1", "Frames and triangulation", "How does diagonal bracing change lateral movement?")
    doc.add_picture(str(ASSETS / "triangulation.png"), width=Inches(5.75))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    callout(doc, "Controlled test", "Build the same two-storey rectangular frame three times: A unbraced, B one diagonal per storey, C X-braced. Use Timber throughout, a 500 kg roof payload, amplitude 30 mm, frequency 1.50 Hz, duration 8 s, damping 8%. Place S1 on the roof and S2 on the first floor.", PALE_GOLD)
    p = doc.add_paragraph("Prediction: Which design will have the lowest roof response? Why?")
    p.runs[0].bold = True
    lines(doc, 1)
    table(doc, ["Design", "Cost ($)", "Max |S1| roof (mm)", "Max |S2| floor (mm)", "Max drift (mm)", "Failures", "Survived?"], [
        ["A Unbraced", "", "", "", "", "", ""],
        ["B Diagonal", "", "", "", "", "", ""],
        ["C X-braced", "", "", "", "", "", ""],
    ], [1.05, 0.72, 1.08, 1.08, 0.86, 0.72, 0.8], 7.7, [0.42]*3)
    question(doc, "Q1", "Make a claim about triangulation. Support it with two numerical results.", 2)
    question(doc, "Q2", "What trade-off appears when more bracing is added? Consider cost, material and access/openings.", 1)

    # Activity 2
    page(doc)
    add_title(doc, "Investigation 2", "Materials: performance and cost", "Does the most expensive material always give the best engineering solution?")
    callout(doc, "Controlled test", "Use your X-braced design from Investigation 1. Repeat the identical shake with every structural member changed to one material at a time. Keep geometry, sensors, 500 kg payload and shake settings unchanged.", PALE_GOLD)
    table(doc, ["Material", "App cost / m", "Strength", "Stiffness", "Weight", "Damping"], [
        ["Timber", "$10", "medium", "medium", "low", "medium"],
        ["Steel", "$30", "high", "high", "high", "low"],
        ["Reinforced", "$22", "high", "high", "medium", "medium"],
        ["Lightweight", "$7", "low", "low", "low", "medium–high"],
    ], [1.25, 1.05, 1.0, 1.0, 1.0, 1.1], 8.3)
    table(doc, ["Material", "Total cost ($)", "Max |S1| (mm)", "Max drift (mm)", "Failures", "Survived?"], [
        ["Timber", "", "", "", "", ""], ["Steel", "", "", "", "", ""],
        ["Reinforced", "", "", "", "", ""], ["Lightweight", "", "", "", "", ""],
    ], [1.25, 1.15, 1.15, 1.15, 0.9, 1.1], 8.3, [0.38]*4)
    question(doc, "Q3", "Which material gave the lowest movement? Which gave the best performance for cost? Use data.", 3)
    question(doc, "Q4", "Explain why material selection is a multi-criteria decision rather than a simple ‘strongest wins’ choice.", 2)

    # Activity 3
    page(doc)
    add_title(doc, "Investigation 3", "Mass and where it is placed", "How does payload position affect movement and structural demand?")
    callout(doc, "Controlled test", "Use one unchanged, two-storey X-braced design. Use Timber and the Investigation 1 shake settings. Compare 2000 kg at the first floor, 2000 kg at the roof, and 1000 kg at each level. Sensors remain at roof (S1) and first floor (S2).", PALE_GOLD)
    table(doc, ["Trial", "Mass arrangement", "Max |S1| roof (mm)", "Max |S2| floor (mm)", "Max drift (mm)", "Failures", "Survived?"], [
        ["A", "2000 kg low", "", "", "", "", ""],
        ["B", "2000 kg high", "", "", "", "", ""],
        ["C", "1000 + 1000 kg", "", "", "", "", ""],
    ], [0.45, 1.35, 1.12, 1.12, 0.95, 0.75, 0.78], 7.8, [0.44]*3)
    question(doc, "Q5", "Which mass arrangement produced the greatest roof response or drift? Quote measurements.", 2)
    question(doc, "Q6", "Use inertia and centre of mass to explain the pattern. Why can heavy upper levels be challenging?", 3)
    question(doc, "Q7", "Was total mass controlled in all three trials? Was mass position controlled? Identify the independent variable.", 2)

    # Activity 4
    page(doc)
    add_title(doc, "Investigation 4", "Frequency, natural period and resonance", "When can repeated shaking build a large response?")
    doc.add_picture(str(ASSETS / "resonance.png"), width=Inches(5.75))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p = doc.add_paragraph("Every structure has natural modes of vibration. When forcing frequency is near a natural frequency, repeated pushes can add efficiently and response may grow. This is resonance. Height, stiffness, mass and damping all influence the response.")
    callout(doc, "Frequency sweep", "Use one unchanged design. Set amplitude 10 mm, duration 8 s and damping 8%. Test the listed frequencies. Record the app’s calculated peak acceleration and the maximum absolute roof-sensor response.", PALE_GOLD)
    table(doc, ["Frequency (Hz)", "0.50", "1.00", "1.50", "2.00", "3.00", "4.00"], [
        ["Peak acceleration (g)", "", "", "", "", "", ""],
        ["Max |S1| (mm)", "", "", "", "", "", ""],
    ], [1.5, 0.83, 0.83, 0.83, 0.83, 0.83, 0.83], 8.2, [0.42, 0.42])
    question(doc, "Q8", "At which frequency was the response largest? Does this suggest a possible resonant region?", 1)
    question(doc, "Q9", "Why is constant amplitude not a perfectly fair resonance test? Hint: compare peak acceleration as frequency rises.", 2)
    question(doc, "Q10", "Predict how added damping would change the height and width of a resonance peak. Test one frequency if time permits.", 1)

    # Open challenge brief
    page(doc)
    add_title(doc, "Open-ended design task", "Resilient Community Hub Challenge", "Design a three-level community building that protects occupants while meeting constraints.")
    callout(doc, "Scenario", "Your town needs a compact community hub that may shelter people after an earthquake. The building needs useful floor space, must carry emergency supplies, and must be affordable. You will independently propose, test and justify a structural concept.", PALE_GOLD)
    doc.add_heading("Non-negotiable constraints", 2)
    table(doc, ["Requirement", "Constraint"], [
        ["Form", "Three occupied levels, fixed to the shake table with Floor Joints"],
        ["Payload", "2000 kg total, placed in 100 kg intervals; justify its distribution"],
        ["Materials", "Use at least two app materials; total app cost ≤ $750"],
        ["Complexity", "Maximum 24 structural members (beams + braces)"],
        ["Measurement", "Use four sensors: base/low, middle, upper and roof locations"],
        ["Research", "Use at least two earthquake-resistant principles and cite sources"],
        ["Final test", "60 mm amplitude • 2.00 Hz • 15 s • 8% damping"],
        ["Performance", "App survival + zero member failures + max roof |Δx| ≤ 120 mm"],
    ], [1.35, 5.35], 8.6)
    doc.add_heading("Engineering process", 2)
    p = doc.add_paragraph("1  DEFINE  →  2  IMAGINE  →  3  PLAN  →  4  BUILD  →  5  TEST  →  6  ANALYSE  →  7  IMPROVE  →  8  COMMUNICATE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    callout(doc, "Open-ended means", "There is no single correct structure. A strong solution is one you can defend with measurements, structural reasoning, research, trade-offs and honest discussion of limitations.", PALE_TEAL)

    # Planning canvas
    page(doc)
    add_title(doc, "Challenge planning", "Concept sketch and design reasoning", "Sketch clearly enough that your teacher or another student could reproduce your idea.")
    t = table(doc, ["Front/elevation sketch — label floor joints, beams, braces, materials, masses and S1–S4", "Key dimensions / notes"], [["", ""]], [5.1, 1.6], 8.5, [2.7])
    for c in t.rows[1].cells:
        shade(c, "FAFBFC")
    table(doc, ["Decision", "My choice", "Reason / research link"], [
        ["Structural form and load path", "", ""],
        ["Bracing / lateral system", "", ""],
        ["Materials", "", ""],
        ["Mass distribution", "", ""],
        ["Sensor locations", "", ""],
    ], [1.7, 1.7, 3.3], 8.3, [0.5]*5)
    question(doc, "C1", "Predict the part or storey most likely to experience the greatest movement or failure. Why?", 1)

    # Iteration data
    page(doc)
    add_title(doc, "Challenge testing", "Test, analyse and improve", "Use a lower-risk diagnostic test before the final hazard test.")
    callout(doc, "Diagnostic test", "Suggested settings: 30 mm amplitude, 1.50 Hz, 8 s, 8% damping. Observe load paths and sensor differences. Change one major feature at a time so you can explain the effect.", PALE_GOLD)
    table(doc, ["Version", "Change made", "Cost ($)", "Max |S1|", "Max |S2|", "Max |S3|", "Max |S4|", "Drift", "Failures"], [
        ["V1", "baseline", "", "", "", "", "", "", ""],
        ["V2", "", "", "", "", "", "", "", ""],
        ["V3", "", "", "", "", "", "", "", ""],
    ], [0.48, 1.55, 0.66, 0.66, 0.66, 0.66, 0.66, 0.6, 0.68], 7.1, [0.45]*3)
    question(doc, "C2", "Which change had the strongest evidence of improvement? Quote before-and-after data.", 3)
    question(doc, "C3", "Did any change improve one measure but worsen another? Explain the trade-off.", 3)
    doc.add_heading("Final hazard test", 2)
    table(doc, ["Cost", "Roof max |Δx|", "Max drift", "Failed members", "Height retained / app result", "Challenge met?"], [["", "", "", "", "", ""]], [0.9, 1.15, 1.05, 1.05, 1.6, 0.95], 8.1, [0.52])
    p = doc.add_paragraph("Attach or paste graph/screenshot evidence here if your teacher requests it.")
    p.runs[0].italic = True
    p.runs[0].font.color.rgb = RGBColor.from_string(MID)

    # Final report
    page(doc)
    add_title(doc, "Communicate", "Engineering evidence report", "Answer in complete sentences. Use precise measurements and acknowledge limitations.")
    question(doc, "E1 — Claim", "How well did your final design meet the brief? State a clear overall judgement.", 3)
    question(doc, "E2 — Evidence", "Give at least three relevant results, with units, including sensor response, drift/failure and cost.", 4)
    question(doc, "E3 — Reasoning", "Explain how two structural features caused the measured response. Use terms such as load path, stiffness, ductility, damping, inertia, triangulation or resonance correctly.", 5)
    question(doc, "E4 — Research connection", "How does one credible real-world source support or challenge a choice you made? Cite it.", 3)
    question(doc, "E5 — Limitations", "Identify two limitations of the simulator or your method. How could a real engineer obtain stronger evidence?", 4)
    question(doc, "E6 — Next iteration", "If given one more design cycle, what specific change would you test and what result would you expect?", 3)

    # Rubric
    page(doc)
    add_title(doc, "Assessment", "Student checklist and rubric", "Use this before submitting your workbook and design evidence.")
    table(doc, ["Criterion", "High achievement", "Sound achievement", "Developing", "Marks"], [
        ["Investigation method", "Fair tests; variables and sensor positions clear; data complete with units", "Mostly controlled tests; usable data", "Variables unclear or data incomplete", "/4"],
        ["Structural understanding", "Accurate, connected explanation of forces and ≥3 principles", "Explains main principles with minor gaps", "Mostly describes results without mechanism", "/4"],
        ["Data analysis", "Patterns, anomalies and trade-offs supported by quantitative evidence", "Uses relevant data to support conclusions", "Limited or unsupported conclusions", "/4"],
        ["Research literacy", "Two credible sources evaluated, paraphrased and cited; limitations recognised", "Relevant sources cited and mostly explained", "Sources weak, copied or not connected", "/4"],
        ["Design process", "Constraints met; purposeful iterations; final decision strongly justified", "Design tested and improved; most constraints met", "Little iteration or weak justification", "/4"],
    ], [1.25, 2.0, 1.72, 1.4, 0.45], 7.3, [0.62]*5)
    p = doc.add_paragraph("TOTAL: ____ / 20")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(12)
    doc.add_heading("Submission checklist", 2)
    for item in [
        "All four investigations include predictions, measurements and answers.",
        "My cited research is in my own words and URLs are recorded.",
        "My challenge design is labelled and meets the constraints, or exceptions are explained.",
        "I included evidence from more than one design version.",
        "My final report uses numerical evidence and discusses limitations.",
    ]:
        doc.add_paragraph("☐ " + item)
    table(doc, ["My strongest design decision", "My most useful next improvement"], [["", ""]], [3.35, 3.35], 8.2, [0.7])

    # Sources
    page(doc)
    add_title(doc, "Research sources", "Trusted starting points", "Use these to begin; your teacher may approve additional sources. Accessed 27 August 2026.")
    add_source(doc, "Earthquakes", "Geoscience Australia", "https://www.ga.gov.au/education/natural-hazards/earthquake", "Australian earthquake processes, magnitude and intensity, local factors, impacts and case examples.")
    add_source(doc, "National Seismic Hazard Assessment", "Geoscience Australia", "https://www.ga.gov.au/scientific-topics/community-safety/data-and-products/nsha", "How national earthquake-hazard information supports risk reduction and building-code decisions.")
    add_source(doc, "Resilience to the shake", "Geoscience Australia", "https://www.ga.gov.au/news/resilience-to-the-shake", "Australian work on the vulnerability and strengthening of older masonry buildings, including Adelaide context.")
    add_source(doc, "Earthquake-Resistant Design Concepts (FEMA P-749)", "FEMA / NEHRP", "https://www.fema.gov/sites/default/files/documents/fema_p-749-earthquake-resistant-design-concepts_112022.pdf", "Accessible technical explanations of strength, stiffness, ductility, bracing, isolation and energy dissipation.")
    add_source(doc, "What are the Effects of Earthquakes?", "U.S. Geological Survey", "https://www.usgs.gov/programs/earthquake-hazards/what-are-effects-earthquakes", "How amplitude, frequency, duration, distance and local ground conditions influence shaking and damage.")
    add_source(doc, "How Seismic Waves Affect Different Size Buildings", "U.S. Geological Survey", "https://www.usgs.gov/programs/earthquake-hazards/how-seismic-waves-affect-different-size-buildings", "Why different building sizes respond differently to wave frequencies.")
    add_source(doc, "Building resonance", "EarthScope Consortium (IRIS)", "https://www.iris.edu/hq/inclass/animation/building_resonance_the_resonant_frequency_of_different_seismic_waves", "Natural period, forcing frequency and resonance explained through building response.")
    callout(doc, "Source use", "Do not copy sentences. Read, close the source, write the idea in your own words, reopen it to check accuracy, then record the organisation, title and URL.", PALE_TEAL)

    header_footer(doc)
    for t in doc.tables:
        set_repeat_table_layout(t)
    core = doc.core_properties
    core.title = "Year 9 Earthquake Engineering Student Task"
    core.subject = "Structural basics investigations and an open-ended earthquake-resistant design challenge"
    core.author = "QuakeLab classroom resource"
    core.keywords = "Year 9, earthquake engineering, structures, STEM, investigation, QuakeLab"
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build_doc()
